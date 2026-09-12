import pdfplumber
from docx import Document
from pathlib import Path

out = Path(r"C:\Users\89722\Documents\ChatGPT\B题思路详解\extracted")
out.mkdir(exist_ok=True)

pdf_map = {
    "B题": r"C:\Users\89722\AppData\Local\Temp\B题 (5).pdf",
    "交汇": r"D:\Wexin\xwechat_files\wxid_5ra4ujb3l8p322_635c\msg\file\2026-09\交汇.pdf",
    "问题2": r"D:\Wexin\xwechat_files\wxid_5ra4ujb3l8p322_635c\msg\file\2026-09\问题2.pdf",
}

for name, f in pdf_map.items():
    parts = []
    with pdfplumber.open(f) as pdf:
        parts.append(f"PAGES: {len(pdf.pages)}\n")
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            parts.append(f"\n----- PAGE {i+1} -----\n{text}")
    (out / f"{name}.txt").write_text("".join(parts), encoding="utf-8")
    print(name, "written", (out / f"{name}.txt").stat().st_size)

doc_map = {
    "附件1": r"C:\Users\89722\AppData\Local\Temp\附件1 (5).docx",
    "附件2": r"C:\Users\89722\AppData\Local\Temp\附件2 (5).docx",
}

for name, f in doc_map.items():
    parts = []
    doc = Document(f)
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text + "\n")
    for ti, table in enumerate(doc.tables):
        parts.append(f"\n[TABLE {ti+1}]\n")
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells) + "\n")
    (out / f"{name}.txt").write_text("".join(parts), encoding="utf-8")
    print(name, "written", (out / f"{name}.txt").stat().st_size)
